import io
import unittest

import docx

import fixtures
from engine.resolver import StyleResolver


def _load(blob):
    document = docx.Document(io.BytesIO(blob))
    return document, StyleResolver(document)


def _first(document, style_name):
    for paragraph in document.paragraphs:
        if paragraph.style.name == style_name:
            return paragraph
    raise AssertionError('không có đoạn nào dùng style %s' % style_name)


class TestRunProps(unittest.TestCase):
    def test_font_thua_huong_tu_style(self):
        """Bẫy chính: run không khai font, phải ra 'Arial' của style."""
        document, resolver = _load(fixtures.sai_font())
        paragraph = _first(document, 'VB_NoiDung')
        run = paragraph.runs[0]
        self.assertIsNone(run.font.name, 'fixture sai: run không được khai font')
        props = resolver.run_props(run, paragraph)
        self.assertEqual(props['font'], 'Arial')

    def test_co_chu_thua_huong_tu_style(self):
        document, resolver = _load(fixtures.chuan_66())
        paragraph = _first(document, 'VB_TieuDeDang')
        props = resolver.run_props(paragraph.runs[0], paragraph)
        self.assertEqual(props['size_pt'], 15.0)
        self.assertEqual(props['font'], 'Times New Roman')

    def test_bold_thua_huong_tu_style(self):
        document, resolver = _load(fixtures.chuan_66())
        dam = resolver.run_props(*_run_and_para(document, 'VB_TrichYeu'))
        thuong = resolver.run_props(*_run_and_para(document, 'VB_NoiDung'))
        self.assertTrue(dam['bold'])
        self.assertFalse(thuong['bold'])

    def test_run_khai_truc_tiep_thang_style(self):
        document, resolver = _load(fixtures.run_lech_nhau())
        paragraph = document.paragraphs[-1]
        self.assertEqual(len(paragraph.runs), 2)
        thua_huong = resolver.run_props(paragraph.runs[0], paragraph)
        khai_truc_tiep = resolver.run_props(paragraph.runs[1], paragraph)
        self.assertEqual(thua_huong['size_pt'], 14.0)
        self.assertEqual(khai_truc_tiep['size_pt'], 11.0)

    def test_run_none_van_ra_dinh_dang_cua_style(self):
        document, resolver = _load(fixtures.chuan_66())
        paragraph = _first(document, 'VB_NoiDung')
        props = resolver.run_props(None, paragraph)
        self.assertEqual(props['size_pt'], 14.0)

    def test_khong_co_style_thi_lay_tu_Normal(self):
        document, resolver = _load(fixtures.khong_co_style())
        paragraph = document.paragraphs[0]
        props = resolver.run_props(paragraph.runs[0], paragraph)
        self.assertEqual(props['font'], 'Times New Roman')
        self.assertEqual(props['size_pt'], 14.0)


class TestParaProps(unittest.TestCase):
    def test_can_le_thua_huong_tu_style(self):
        document, resolver = _load(fixtures.chuan_66())
        self.assertEqual(
            resolver.para_props(_first(document, 'VB_NoiDung'), 14.0)['align'],
            'justify')
        self.assertEqual(
            resolver.para_props(_first(document, 'VB_TieuDeDang'), 15.0)['align'],
            'center')
        self.assertEqual(
            resolver.para_props(_first(document, 'VB_ChuKy'), 14.0)['align'],
            'right')

    def test_dan_dong_auto(self):
        document, resolver = _load(fixtures.chuan_66())
        props = resolver.para_props(_first(document, 'VB_NoiDung'), 14.0)
        self.assertAlmostEqual(props['line_spacing'], 1.5, places=3)
        self.assertFalse(props['line_spacing_fixed'])

    def test_dan_dong_exact_danh_dau_fixed(self):
        document, resolver = _load(fixtures.sai_dan_dong_exact())
        props = resolver.para_props(_first(document, 'VB_NoiDung'), 14.0)
        self.assertTrue(props['line_spacing_fixed'])

    def test_thut_dau_dong_thua_huong_tu_style(self):
        document, resolver = _load(fixtures.chuan_66())
        props = resolver.para_props(_first(document, 'VB_NoiDung'), 14.0)
        self.assertAlmostEqual(props['first_line_indent_cm'], 1.27, places=2)


def _run_and_para(document, style_name):
    paragraph = _first(document, style_name)
    return paragraph.runs[0], paragraph


if __name__ == '__main__':
    unittest.main()
