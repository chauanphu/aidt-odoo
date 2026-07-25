import io
import unittest
import zipfile

import docx
from docx.opc.exceptions import PackageNotFoundError

import fixtures

DOCX_HOP_LE = (
    'chuan_66', 'chuan_nd30', 'sai_font', 'sai_dan_dong_exact', 'sai_le_trang',
    'thieu_noi_nhan', 'run_lech_nhau', 'in_dam_giua_cau', 'khong_co_style',
)


class TestFixtures(unittest.TestCase):
    def test_moi_fixture_mo_duoc_bang_python_docx(self):
        for name in DOCX_HOP_LE:
            with self.subTest(fixture=name):
                blob = getattr(fixtures, name)()
                document = docx.Document(io.BytesIO(blob))
                self.assertGreater(len(document.paragraphs), 0)

    def test_hong_khong_mo_duoc(self):
        # python-docx 1.1.2 chỉ dịch lỗi sang PackageNotFoundError khi nhận
        # ĐƯỜNG DẪN; với stream (BytesIO — cách engine luôn dùng) nó để
        # zipfile.BadZipFile lọt thẳng ra. Chấp nhận cả hai để test không vỡ
        # khi đổi phiên bản thư viện. Hợp đồng thật là parse_docx phải ném
        # UnreadableDocx, và Task 4 kiểm điều đó.
        with self.assertRaises((PackageNotFoundError, zipfile.BadZipFile)):
            docx.Document(io.BytesIO(fixtures.hong()))

    def test_chuan_nd30_dung_co_chu_cua_ND30(self):
        """chuan_nd30 phải đạt theo ND-30 thật, không dùng cỡ của chuẩn Đảng."""
        document = docx.Document(io.BytesIO(fixtures.chuan_nd30()))
        self.assertEqual(document.styles['VB_SoKyHieu'].font.size.pt, 13)
        self.assertEqual(document.styles['VB_NoiNhan'].font.size.pt, 11)

    def test_sai_font_khai_o_style_khong_khai_o_run(self):
        """Chốt cái bẫy: run.font.name là None, còn style mới giữ 'Arial'."""
        document = docx.Document(io.BytesIO(fixtures.sai_font()))
        noi_dung = [p for p in document.paragraphs if p.style.name == 'VB_NoiDung']
        self.assertTrue(noi_dung)
        for run in noi_dung[0].runs:
            self.assertIsNone(run.font.name)
        self.assertEqual(document.styles['VB_NoiDung'].font.name, 'Arial')

    def test_thieu_noi_nhan_thi_thieu_that(self):
        document = docx.Document(io.BytesIO(fixtures.thieu_noi_nhan()))
        self.assertFalse(
            [p for p in document.paragraphs if p.style.name == 'VB_NoiNhan'])

    def test_khong_co_style_khong_dung_VB(self):
        document = docx.Document(io.BytesIO(fixtures.khong_co_style()))
        for paragraph in document.paragraphs:
            self.assertFalse(paragraph.style.name.startswith('VB_'))
