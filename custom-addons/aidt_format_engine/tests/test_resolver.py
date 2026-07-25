import io
import unittest

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Pt

import fixtures
from aidt_format_engine.resolver import StyleResolver


def _load(blob):
    document = docx.Document(io.BytesIO(blob))
    return document, StyleResolver(document)


def _first(document, style_name):
    for paragraph in document.paragraphs:
        if paragraph.style.name == style_name:
            return paragraph
    raise AssertionError('không có đoạn nào dùng style %s' % style_name)


class TestRunProps(unittest.TestCase):
    def test_font_thua_huong_tu_style(self):
        """Bẫy chính: run không khai font, phải ra 'Arial' của style."""
        document, resolver = _load(fixtures.sai_font())
        paragraph = _first(document, 'VB_NoiDung')
        run = paragraph.runs[0]
        self.assertIsNone(run.font.name, 'fixture sai: run không được khai font')
        props = resolver.run_props(run, paragraph)
        self.assertEqual(props['font'], 'Arial')

    def test_co_chu_thua_huong_tu_style(self):
        document, resolver = _load(fixtures.chuan_66())
        paragraph = _first(document, 'VB_TieuDeDang')
        props = resolver.run_props(paragraph.runs[0], paragraph)
        self.assertEqual(props['size_pt'], 15.0)
        self.assertEqual(props['font'], 'Times New Roman')

    def test_bold_thua_huong_tu_style(self):
        document, resolver = _load(fixtures.chuan_66())
        dam = resolver.run_props(*_run_and_para(document, 'VB_TrichYeu'))
        thuong = resolver.run_props(*_run_and_para(document, 'VB_NoiDung'))
        self.assertTrue(dam['bold'])
        self.assertFalse(thuong['bold'])

    def test_run_khai_truc_tiep_thang_style(self):
        document, resolver = _load(fixtures.run_lech_nhau())
        paragraph = document.paragraphs[-1]
        self.assertEqual(len(paragraph.runs), 2)
        thua_huong = resolver.run_props(paragraph.runs[0], paragraph)
        khai_truc_tiep = resolver.run_props(paragraph.runs[1], paragraph)
        self.assertEqual(thua_huong['size_pt'], 14.0)
        self.assertEqual(khai_truc_tiep['size_pt'], 11.0)
        self.assertEqual(khai_truc_tiep['font'], 'Times New Roman')
        self.assertFalse(khai_truc_tiep['bold'])

    def test_run_none_van_ra_dinh_dang_cua_style(self):
        document, resolver = _load(fixtures.chuan_66())
        paragraph = _first(document, 'VB_NoiDung')
        props = resolver.run_props(None, paragraph)
        self.assertEqual(props['size_pt'], 14.0)

    def test_khong_co_style_thi_lay_tu_Normal(self):
        document, resolver = _load(fixtures.khong_co_style())
        paragraph = document.paragraphs[0]
        props = resolver.run_props(paragraph.runs[0], paragraph)
        self.assertEqual(props['font'], 'Times New Roman')
        self.assertEqual(props['size_pt'], 14.0)

    def test_font_giai_duoc_qua_theme_khi_khong_tang_nao_khai(self):
        """docx trần không khai font ở đâu: phải giải ra font theme, không ra None."""
        document = docx.Document()
        paragraph = document.add_paragraph('Nội dung không khai font')
        resolver = StyleResolver(document)
        props = resolver.run_props(paragraph.runs[0], paragraph)
        self.assertIsNotNone(props['font'],
                             'font phải giải được qua theme, không được là None')

    def test_character_style_thang_paragraph_style(self):
        """Character style nằm trên paragraph style trong thứ tự ưu tiên."""
        document = docx.Document()
        para_style = document.styles.add_style('P_Test', WD_STYLE_TYPE.PARAGRAPH)
        para_style.font.name = 'Times New Roman'
        para_style.font.size = Pt(14)
        char_style = document.styles.add_style('C_Test', WD_STYLE_TYPE.CHARACTER)
        char_style.font.name = 'Arial'

        paragraph = document.add_paragraph(style='P_Test')
        run = paragraph.add_run('chữ dùng character style')
        run.style = char_style

        props = StyleResolver(document).run_props(run, paragraph)
        self.assertEqual(props['font'], 'Arial', 'character style phải thắng')
        self.assertEqual(props['size_pt'], 14.0,
                         'cỡ chữ không khai ở character style thì giữ của paragraph style')

    def test_leo_chuoi_base_style_va_cache_khong_lan(self):
        """Style con thừa hưởng từ base, và cache không làm hỏng style base."""
        document = docx.Document()
        base = document.styles.add_style('VB_Base', WD_STYLE_TYPE.PARAGRAPH)
        base.font.name = 'Times New Roman'
        base.font.size = Pt(14)
        base.font.bold = False
        con = document.styles.add_style('VB_Con', WD_STYLE_TYPE.PARAGRAPH)
        con.base_style = base
        con.font.bold = True

        p_con = document.add_paragraph('đoạn dùng style con', style='VB_Con')
        p_base = document.add_paragraph('đoạn dùng style base', style='VB_Base')
        resolver = StyleResolver(document)

        props_con = resolver.run_props(p_con.runs[0], p_con)
        self.assertEqual(props_con['font'], 'Times New Roman',
                         'font phải thừa hưởng từ base_style')
        self.assertEqual(props_con['size_pt'], 14.0)
        self.assertTrue(props_con['bold'], 'style con ghi đè bold')

        # Giải style con TRƯỚC rồi mới tới base: nếu cache bị lẫn, base sẽ
        # dính bold=True của con.
        props_base = resolver.run_props(p_base.runs[0], p_base)
        self.assertFalse(props_base['bold'],
                         'cache không được để style con làm bẩn style base')

    def test_font_khai_o_hAnsi_ma_khong_co_ascii(self):
        """File Word thật có thể chỉ khai hAnsi — slot chi phối chữ tiếng Việt."""
        document = docx.Document()
        paragraph = document.add_paragraph()
        run = paragraph.add_run('văn bản có dấu tiếng Việt')
        rpr = run._r.get_or_add_rPr()
        fonts = rpr.get_or_add_rFonts()
        fonts.set(qn('w:hAnsi'), 'Times New Roman')   # cố ý KHÔNG đặt w:ascii

        props = StyleResolver(document).run_props(run, paragraph)
        self.assertEqual(props['font'], 'Times New Roman')


class TestParaProps(unittest.TestCase):
    def test_can_le_thua_huong_tu_style(self):
        document, resolver = _load(fixtures.chuan_66())
        self.assertEqual(
            resolver.para_props(_first(document, 'VB_NoiDung'), 14.0)['align'],
            'justify')
        self.assertEqual(
            resolver.para_props(_first(document, 'VB_TieuDeDang'), 15.0)['align'],
            'center')
        self.assertEqual(
            resolver.para_props(_first(document, 'VB_ChuKy'), 14.0)['align'],
            'right')

    def test_dan_dong_auto(self):
        document, resolver = _load(fixtures.chuan_66())
        props = resolver.para_props(_first(document, 'VB_NoiDung'), 14.0)
        self.assertAlmostEqual(props['line_spacing'], 1.5, places=3)
        self.assertFalse(props['line_spacing_fixed'])

    def test_dan_dong_exact_danh_dau_fixed(self):
        document, resolver = _load(fixtures.sai_dan_dong_exact())
        props = resolver.para_props(_first(document, 'VB_NoiDung'), 14.0)
        self.assertTrue(props['line_spacing_fixed'])

    def test_thut_dau_dong_thua_huong_tu_style(self):
        document, resolver = _load(fixtures.chuan_66())
        props = resolver.para_props(_first(document, 'VB_NoiDung'), 14.0)
        self.assertAlmostEqual(props['first_line_indent_cm'], 1.27, places=2)


def _run_and_para(document, style_name):
    paragraph = _first(document, style_name)
    return paragraph.runs[0], paragraph


if __name__ == '__main__':
    unittest.main()
