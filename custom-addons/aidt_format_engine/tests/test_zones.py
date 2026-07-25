import io
import unittest

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH

import fixtures
from aidt_format_engine.parser import parse_docx
from aidt_format_engine.zones import ZONES, detect_zones


def _prepare(blob):
    return detect_zones(parse_docx(blob))


def _zones_of(doc):
    return {p.zone for p in doc.paras if p.text.strip()}


def _nhieu_doan(muc):
    """Dựng .docx trần từ danh sách (chữ, căn lề) — không style VB_* nào.

    Dùng cho các test nhận diện theo cấu trúc: thứ cần dựng là VỊ TRÍ của đoạn
    so với các đoạn khác, nên fixture phải nhỏ và đọc được ngay tại chỗ.
    """
    document = docx.Document()
    for text, *align in muc:
        paragraph = document.add_paragraph(text)
        if align and align[0] is not None:
            paragraph.paragraph_format.alignment = align[0]
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _mot_doan(text, align=None):
    return _nhieu_doan([(text, align)])


def _zone_cua(doc, index):
    return [p for p in doc.paras if p.text.strip()][index].zone


class TestZoneTheoStyle(unittest.TestCase):
    def test_moi_vung_nhan_dien_qua_style(self):
        doc = _prepare(fixtures.chuan_66())
        self.assertEqual(
            _zones_of(doc),
            {'tieu_de_dang', 'so_ky_hieu', 'trich_yeu', 'noi_dung',
             'noi_nhan', 'chu_ky'})

    def test_confidence_la_style_cho_moi_doan_co_chu(self):
        doc = _prepare(fixtures.chuan_66())
        for para in doc.paras:
            if para.text.strip():
                self.assertEqual(para.zone_confidence, 'style',
                                 'đoạn %d' % para.index)

    def test_moi_zone_deu_nam_trong_ZONES(self):
        doc = _prepare(fixtures.chuan_66())
        for para in doc.paras:
            self.assertIn(para.zone, ZONES)

    def test_style_thang_ca_khi_heuristic_noi_khac(self):
        """Đoạn mang style VB_NoiNhan nhưng chữ khớp mẫu số ký hiệu.

        Heuristic sẽ đoán 'so_ky_hieu'; style phải thắng. Đây là chỗ duy nhất
        hai đường cho câu trả lời khác nhau, nên là chỗ duy nhất chứng minh
        được đường nào có quyền cao hơn.
        """
        document = docx.Document()
        style = document.styles.add_style('VB_NoiNhan', WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = 'Times New Roman'
        document.add_paragraph('Số: 99-CV/TU', style='VB_NoiNhan')

        buffer = io.BytesIO()
        document.save(buffer)
        doc = detect_zones(parse_docx(buffer.getvalue()))

        khop = [p for p in doc.paras if p.text.strip() == 'Số: 99-CV/TU']
        self.assertEqual(len(khop), 1)
        self.assertEqual(khop[0].zone, 'noi_nhan',
                         'style phải thắng heuristic khi hai bên nói khác nhau')
        self.assertEqual(khop[0].zone_confidence, 'style')


class TestZoneHeuristic(unittest.TestCase):
    def test_gan_dung_khi_khong_co_style(self):
        doc = _prepare(fixtures.khong_co_style())
        by_zone = {}
        for para in doc.paras:
            if para.text.strip():
                by_zone.setdefault(para.zone, []).append(para.text)
        self.assertIn('tieu_de_dang', by_zone)
        self.assertIn('so_ky_hieu', by_zone)
        self.assertIn('trich_yeu', by_zone)
        self.assertIn('noi_nhan', by_zone)
        self.assertIn('chu_ky', by_zone)

    def test_confidence_la_heuristic(self):
        doc = _prepare(fixtures.khong_co_style())
        for para in doc.paras:
            if para.text.strip():
                self.assertEqual(para.zone_confidence, 'heuristic')

    def test_so_ky_hieu_khop_ca_dau_gach_va_gach_cheo(self):
        doc = _prepare(fixtures.khong_co_style())
        so = [p for p in doc.paras if p.zone == 'so_ky_hieu']
        self.assertEqual(len(so), 1)
        self.assertIn('456-BC/TU', so[0].text)


class TestZoneTrongBang(unittest.TestCase):
    def test_tieu_de_dang_va_so_ky_hieu_nhan_dien_du_nam_trong_bang(self):
        """C1: khối đầu trang dựng bằng bảng — style vẫn phải thắng, dù đoạn
        nằm trong ô bảng chứ không phải paragraph rời của body."""
        doc = _prepare(fixtures.vb_that_dau_trang_bang())
        zones = _zones_of(doc)
        self.assertIn('tieu_de_dang', zones)
        self.assertIn('so_ky_hieu', zones)


class TestTrichYeuVeViec(unittest.TestCase):
    def test_ve_viec_nhan_dien_la_trich_yeu(self):
        """C2: 'Về việc ...' (văn bản có tên loại) phải được heuristic nhận
        là trich_yeu, y hệt 'V/v ...' (công văn)."""
        doc = _prepare(fixtures.vb_that_co_ten_loai())
        self.assertIn('trich_yeu', _zones_of(doc))
        trich_yeu = [p for p in doc.paras if p.zone == 'trich_yeu']
        self.assertEqual(len(trich_yeu), 1)
        self.assertTrue(trich_yeu[0].text.startswith('Về việc'))
        self.assertEqual(trich_yeu[0].zone_confidence, 'heuristic')


class TestTieuDeDangKhongTuMauThuan(unittest.TestCase):
    def test_can_phai_khong_con_duoc_nhan_la_tieu_de_dang(self):
        """I6: heuristic từng nhận align center HOẶC right cho tiêu đề Đảng,
        trong khi bộ luật chỉ cho phép center — tự phát hiện rồi tự phạt.
        Sau khi thu hẹp điều kiện, đoạn căn phải không còn được gán zone
        này (rơi về noi_dung mặc định), nên không còn tự mâu thuẫn."""
        document = docx.Document()
        paragraph = document.add_paragraph(fixtures.TIEU_DE_DANG)
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        buffer = io.BytesIO()
        document.save(buffer)
        doc = detect_zones(parse_docx(buffer.getvalue()))

        khop = [p for p in doc.paras if p.text == fixtures.TIEU_DE_DANG][0]
        self.assertNotEqual(khop.zone, 'tieu_de_dang')

    def test_can_giua_van_duoc_nhan_la_tieu_de_dang(self):
        """Đối chứng: center vẫn phải nhận diện được như trước."""
        document = docx.Document()
        paragraph = document.add_paragraph(fixtures.TIEU_DE_DANG)
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

        buffer = io.BytesIO()
        document.save(buffer)
        doc = detect_zones(parse_docx(buffer.getvalue()))

        khop = [p for p in doc.paras if p.text == fixtures.TIEU_DE_DANG][0]
        self.assertEqual(khop.zone, 'tieu_de_dang')


# -- I4: năm vùng thể thức bổ sung ------------------------------------------
#
# Trước I4, tiêu ngữ / tên cơ quan ban hành / địa danh-ngày tháng / tên loại
# văn bản / họ tên người ký đều rơi vào 'noi_dung' và bị đo bằng thước của
# phần thân (căn đều hai bên, thụt đầu dòng 1cm...) — sai hoàn toàn về bản
# chất, và mỗi vùng sinh vài cảnh báo không phải vi phạm thật.

class TestNamVungBoSung(unittest.TestCase):
    def test_du_muoi_hai_vung_qua_style(self):
        doc = _prepare(fixtures.vb_day_du_theo_style())
        self.assertEqual([p.zone for p in doc.paras if p.text.strip()],
                         fixtures.VUNG_DAY_DU)

    def test_du_muoi_hai_vung_qua_heuristic(self):
        """Hai đường phải cho cùng kết quả — nếu không, heuristic vô dụng đúng
        ở chỗ cần nó nhất: văn bản soạn tay ngoài mẫu."""
        doc = _prepare(fixtures.vb_day_du_khong_style())
        self.assertEqual([p.zone for p in doc.paras if p.text.strip()],
                         fixtures.VUNG_DAY_DU)

    def test_moi_vung_moi_deu_nam_trong_ZONES(self):
        for zone in ('tieu_ngu', 'ten_co_quan', 'dia_danh_ngay', 'ten_loai',
                     'ho_ten_nguoi_ky'):
            self.assertIn(zone, ZONES)


class TestTieuNgu(unittest.TestCase):
    def test_nhan_dien_du_dung_gach_ngang_dai(self):
        """Word tự đổi '-' thành '–' khi gõ. Cả hai đều là tiêu ngữ hợp lệ."""
        doc = _prepare(_mot_doan('Độc lập – Tự do – Hạnh phúc',
                                 WD_ALIGN_PARAGRAPH.CENTER))
        self.assertEqual(_zone_cua(doc, 0), 'tieu_ngu')

    def test_khong_nhan_nham_cau_van_co_chua_cum_tu(self):
        doc = _prepare(_mot_doan(
            'Phát huy tinh thần Độc lập - Tự do - Hạnh phúc trong công tác.',
            WD_ALIGN_PARAGRAPH.JUSTIFY))
        self.assertEqual(_zone_cua(doc, 0), 'noi_dung')


class TestDiaDanhNgay(unittest.TestCase):
    def test_nhan_dien_dang_day_du(self):
        doc = _prepare(_mot_doan('Hà Nội, ngày 25 tháng 7 năm 2026'))
        self.assertEqual(_zone_cua(doc, 0), 'dia_danh_ngay')

    def test_nhan_dien_mau_chua_dien_ngay_thang(self):
        """Mẫu in sẵn để trống ngày và tháng, người ký điền tay sau."""
        doc = _prepare(_mot_doan('Bình Dương, ngày     tháng     năm 2026'))
        self.assertEqual(_zone_cua(doc, 0), 'dia_danh_ngay')

    def test_khong_nhan_nham_cau_van_co_ngay_thang(self):
        doc = _prepare(_mot_doan(
            'Các đơn vị báo cáo kết quả, ngày 30 tháng 9 năm 2026 là hạn cuối '
            'để gửi về Văn phòng theo quy định.'))
        self.assertEqual(_zone_cua(doc, 0), 'noi_dung')


class TestTenCoQuanVaTenLoai(unittest.TestCase):
    """Tên cơ quan và tên loại trông giống hệt nhau — in hoa, đậm, căn giữa.
    Thứ duy nhất phân biệt được là chỗ đứng so với số ký hiệu."""

    def test_cung_mot_chu_in_hoa_tren_so_ky_hieu_la_ten_co_quan(self):
        doc = _prepare(_nhieu_doan([
            ('BÁO CÁO', WD_ALIGN_PARAGRAPH.CENTER),
            ('Số: 12/BC-UBND', WD_ALIGN_PARAGRAPH.CENTER),
            ('V/v kết quả công tác', WD_ALIGN_PARAGRAPH.CENTER),
        ]))
        self.assertEqual(_zone_cua(doc, 0), 'ten_co_quan')

    def test_cung_mot_chu_in_hoa_duoi_so_ky_hieu_la_ten_loai(self):
        doc = _prepare(_nhieu_doan([
            ('Số: 12/BC-UBND', WD_ALIGN_PARAGRAPH.CENTER),
            ('BÁO CÁO', WD_ALIGN_PARAGRAPH.CENTER),
            ('V/v kết quả công tác', WD_ALIGN_PARAGRAPH.CENTER),
        ]))
        self.assertEqual(_zone_cua(doc, 1), 'ten_loai')

    def test_khong_co_so_ky_hieu_thi_khong_doan_ten_co_quan(self):
        """Không có mốc thì không đoán. Đoán bừa ở đây nghĩa là đem thước của
        vùng tên cơ quan đo một đoạn bất kỳ."""
        doc = _prepare(_nhieu_doan([
            ('BÁO CÁO KẾT QUẢ CÔNG TÁC', WD_ALIGN_PARAGRAPH.CENTER),
            ('Nội dung báo cáo trình bày dưới đây.', WD_ALIGN_PARAGRAPH.JUSTIFY),
        ]))
        self.assertEqual(_zone_cua(doc, 0), 'noi_dung')

    def test_ten_loai_phai_dung_ngay_truoc_trich_yeu(self):
        """Chữ in hoa căn giữa nằm giữa phần thân không phải tên loại —
        văn bản hành chính hay có tiêu đề mục viết hoa."""
        doc = _prepare(_nhieu_doan([
            ('Số: 12/BC-UBND', WD_ALIGN_PARAGRAPH.CENTER),
            ('V/v kết quả công tác', WD_ALIGN_PARAGRAPH.CENTER),
            ('PHẦN THỨ NHẤT', WD_ALIGN_PARAGRAPH.CENTER),
            ('Nội dung phần thứ nhất.', WD_ALIGN_PARAGRAPH.JUSTIFY),
        ]))
        self.assertEqual(_zone_cua(doc, 2), 'noi_dung')


class TestHoTenNguoiKy(unittest.TestCase):
    _THAN = [
        ('Số: 12/BC-UBND', WD_ALIGN_PARAGRAPH.CENTER),
        ('V/v kết quả công tác', WD_ALIGN_PARAGRAPH.CENTER),
        ('Nội dung văn bản trình bày ở đây.', WD_ALIGN_PARAGRAPH.JUSTIFY),
    ]

    def test_ten_rieng_sau_chuc_vu_la_ho_ten_nguoi_ky(self):
        doc = _prepare(_nhieu_doan(self._THAN + [
            ('CHỦ TỊCH', WD_ALIGN_PARAGRAPH.RIGHT),
            ('Nguyễn Văn A', WD_ALIGN_PARAGRAPH.RIGHT),
        ]))
        self.assertEqual(_zone_cua(doc, 3), 'chu_ky')
        self.assertEqual(_zone_cua(doc, 4), 'ho_ten_nguoi_ky')

    def test_khong_co_chuc_vu_dung_truoc_thi_khong_doan(self):
        """Không có chức vụ đứng trên thì một dòng căn phải cuối văn bản có
        thể là bất cứ thứ gì. Đoán ở đây là đem thước vùng chữ ký đo bừa."""
        doc = _prepare(_nhieu_doan(self._THAN + [
            ('Nguyễn Văn A', WD_ALIGN_PARAGRAPH.RIGHT),
        ]))
        self.assertEqual(_zone_cua(doc, 3), 'noi_dung')


class TestDongLietKeNoiNhan(unittest.TestCase):
    def test_dong_liet_ke_duoi_nhan_thuoc_vung_noi_nhan(self):
        """I1: chỉ nhãn 'Nơi nhận:' khớp regex; danh sách cơ quan bên dưới
        trước đây rơi vào vùng nội dung và bị đo bằng thước phần thân."""
        doc = _prepare(_nhieu_doan([
            ('Nơi nhận:', None),
            ('- Như trên;', None),
            ('- Lưu: VT.', None),
        ]))
        self.assertEqual([_zone_cua(doc, i) for i in range(3)],
                         ['noi_nhan', 'noi_nhan', 'noi_nhan'])

    def test_gach_dau_dong_khong_dung_sau_noi_nhan_van_la_noi_dung(self):
        """Văn bản hành chính gạch đầu dòng khắp phần thân. Bỏ điều kiện
        'đoạn liền trước là nơi nhận' thì cả phần thân bị hút vào vùng đó."""
        doc = _prepare(_nhieu_doan([
            ('Các đơn vị triển khai những việc sau:', WD_ALIGN_PARAGRAPH.JUSTIFY),
            ('- Rà soát toàn bộ hồ sơ tồn đọng;', WD_ALIGN_PARAGRAPH.JUSTIFY),
        ]))
        self.assertEqual(_zone_cua(doc, 1), 'noi_dung')


class TestVongHaiKhongPhaGiaoUocStyle(unittest.TestCase):
    def test_style_van_thang_o_vong_nhan_dien_theo_cau_truc(self):
        """Giao ước xuyên suốt: style thắng heuristic. Vòng hai chạy SAU nên
        ghi đè được mọi thứ nếu thiếu guard zone_confidence — đoạn dưới mang
        style VB_NoiDung và thỏa đủ điều kiện cấu trúc của tên cơ quan
        (in hoa, đứng trên số ký hiệu), nên nó là chỗ duy nhất chứng minh
        được guard đó còn nguyên."""
        document = docx.Document()
        style = document.styles.add_style('VB_NoiDung', WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = 'Times New Roman'
        document.add_paragraph('ỦY BAN NHÂN DÂN TỈNH BÌNH DƯƠNG',
                               style='VB_NoiDung')
        document.add_paragraph('Số: 12/BC-UBND')

        buffer = io.BytesIO()
        document.save(buffer)
        doc = detect_zones(parse_docx(buffer.getvalue()))

        self.assertEqual(_zone_cua(doc, 0), 'noi_dung')


class TestDieuKienThuHepCuaVongHai(unittest.TestCase):
    def test_doan_toan_so_khong_thanh_ten_co_quan(self):
        """'145/2026' bằng đúng chính nó viết hoa. Nếu _la_in_hoa() chỉ so
        text == text.upper() thì mọi đoạn toàn số, gạch, dấu chấm đứng trên
        số ký hiệu đều thành tên cơ quan ban hành."""
        doc = _prepare(_nhieu_doan([
            ('145/2026', WD_ALIGN_PARAGRAPH.CENTER),
            ('Số: 12/BC-UBND', WD_ALIGN_PARAGRAPH.CENTER),
        ]))
        self.assertEqual(_zone_cua(doc, 0), 'noi_dung')

    def test_cau_van_dai_can_phai_cuoi_van_ban_khong_thanh_ho_ten(self):
        """Bỏ ngưỡng độ dài thì một câu căn phải bất kỳ đứng sau chữ ký đều
        thành họ tên người ký."""
        doc = _prepare(_nhieu_doan([
            ('Số: 12/BC-UBND', WD_ALIGN_PARAGRAPH.CENTER),
            ('V/v kết quả công tác', WD_ALIGN_PARAGRAPH.CENTER),
            ('Nội dung văn bản trình bày ở đây.', WD_ALIGN_PARAGRAPH.JUSTIFY),
            ('CHỦ TỊCH', WD_ALIGN_PARAGRAPH.RIGHT),
            ('Văn bản này thay thế các văn bản trước đây trái với quy định '
             'nêu trên, các đơn vị nghiêm túc thi hành.',
             WD_ALIGN_PARAGRAPH.RIGHT),
        ]))
        self.assertEqual(_zone_cua(doc, 4), 'noi_dung')

    def test_doan_trong_xen_giua_khong_lam_dut_quan_he_vi_tri(self):
        """Người soạn gõ Enter lấy khoảng cách thay vì đặt spacing — Word để
        lại đoạn trống thật giữa tên loại và trích yếu, giữa nhãn nơi nhận và
        danh sách. Nhìn sang đoạn kề mà không bỏ qua đoạn trống thì cả hai
        quan hệ vị trí đứt, và đúng loại văn bản soạn tay là loại hay có
        đoạn trống nhất."""
        doc = _prepare(_nhieu_doan([
            ('Số: 12/BC-UBND', WD_ALIGN_PARAGRAPH.CENTER),
            ('BÁO CÁO', WD_ALIGN_PARAGRAPH.CENTER),
            ('', None),
            ('V/v kết quả công tác', WD_ALIGN_PARAGRAPH.CENTER),
            ('Nội dung báo cáo trình bày dưới đây.', WD_ALIGN_PARAGRAPH.JUSTIFY),
            ('Nơi nhận:', None),
            ('', None),
            ('- Như trên;', None),
        ]))
        self.assertEqual(_zone_cua(doc, 1), 'ten_loai')
        self.assertEqual(_zone_cua(doc, 5), 'noi_nhan')

    def test_tieu_ngu_phai_dung_mot_minh_tren_mot_dong(self):
        """Neo cuối chuỗi: câu văn MỞ ĐẦU bằng đúng cụm tiêu ngữ rồi viết
        tiếp vẫn là câu văn, không phải vùng tiêu ngữ."""
        doc = _prepare(_mot_doan(
            'Độc lập - Tự do - Hạnh phúc là mục tiêu xuyên suốt của cách mạng.',
            WD_ALIGN_PARAGRAPH.JUSTIFY))
        self.assertEqual(_zone_cua(doc, 0), 'noi_dung')


class TestStandardHint(unittest.TestCase):
    def test_tieu_de_dang_suy_ra_dang(self):
        self.assertEqual(_prepare(fixtures.chuan_66()).standard_hint, 'dang')

    def test_quoc_hieu_suy_ra_hanh_chinh(self):
        self.assertEqual(
            _prepare(fixtures.chuan_nd30()).standard_hint, 'hanh_chinh')


if __name__ == '__main__':
    unittest.main()
