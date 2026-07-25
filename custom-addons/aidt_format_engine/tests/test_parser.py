import io
import unittest

import fixtures
from aidt_format_engine.parser import UnreadableDocx, parse_docx


def _blob_of(document):
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestParser(unittest.TestCase):
    def test_le_trang_doc_dung(self):
        doc = parse_docx(fixtures.chuan_66())
        self.assertAlmostEqual(doc.pages.margin_mm['top'], 22.0, places=1)
        self.assertAlmostEqual(doc.pages.margin_mm['left'], 32.0, places=1)
        self.assertAlmostEqual(doc.pages.width_mm, 210.0, places=1)
        self.assertAlmostEqual(doc.pages.height_mm, 297.0, places=1)

    def test_le_sai_doc_ra_gia_tri_sai(self):
        doc = parse_docx(fixtures.sai_le_trang())
        self.assertAlmostEqual(doc.pages.margin_mm['top'], 10.0, places=1)

    def test_moi_doan_co_style_name_va_dinh_dang(self):
        doc = parse_docx(fixtures.chuan_66())
        self.assertGreater(len(doc.paras), 5)
        noi_dung = [p for p in doc.paras if p.style_name == 'VB_NoiDung']
        self.assertTrue(noi_dung)
        self.assertEqual(noi_dung[0].fmt.font, 'Times New Roman')
        self.assertEqual(noi_dung[0].fmt.size_pt, 14.0)
        self.assertEqual(noi_dung[0].fmt.align, 'justify')

    def test_index_bat_dau_tu_0_va_lien_tuc(self):
        doc = parse_docx(fixtures.chuan_66())
        self.assertEqual([p.index for p in doc.paras],
                         list(range(len(doc.paras))))

    def test_runs_conflict_bat_dung_cho(self):
        doc = parse_docx(fixtures.run_lech_nhau())
        lech = [p for p in doc.paras if p.runs_conflict]
        self.assertEqual(len(lech), 1, 'chỉ đúng một đoạn được đánh dấu lệch')
        self.assertIn('cỡ 11', lech[0].text)

    def test_runs_conflict_khong_bat_oan(self):
        doc = parse_docx(fixtures.chuan_66())
        self.assertEqual([p for p in doc.paras if p.runs_conflict], [])

    def test_dinh_dang_lay_tu_run_dai_nhat(self):
        """Một từ lạc cỡ không được quyết định định dạng của cả đoạn."""
        doc = parse_docx(fixtures.run_lech_nhau())
        lech = [p for p in doc.paras if p.runs_conflict][0]
        self.assertEqual(lech.fmt.size_pt, 14.0)

    def test_standard_hint_chua_duoc_dien_o_tang_parser(self):
        self.assertIsNone(parse_docx(fixtures.chuan_66()).standard_hint)

    def test_zone_chua_duoc_dien_o_tang_parser(self):
        doc = parse_docx(fixtures.chuan_66())
        self.assertEqual({p.zone for p in doc.paras}, {None})

    def test_file_hong_bao_UnreadableDocx(self):
        with self.assertRaises(UnreadableDocx):
            parse_docx(fixtures.hong())

    def test_in_dam_giua_cau_khong_bi_danh_dau_lech(self):
        """bold KHÔNG nằm trong tiêu chí runs_conflict — nhấn mạnh là hợp lệ."""
        doc = parse_docx(fixtures.in_dam_giua_cau())
        khop = [p for p in doc.paras
                if p.text.startswith('Các đơn vị hoàn thành trước')]
        self.assertEqual(len(khop), 1, 'fixture phải có đúng một đoạn như vậy')
        self.assertFalse(khop[0].runs_conflict,
                         'in đậm giữa câu không phải lỗi lệch định dạng')

    def test_le_khong_khai_thi_tra_None_chu_khong_sap(self):
        """Section kế thừa lề: python-docx trả None, parser không được sập."""
        import docx as docx_mod
        from docx.oxml.ns import qn

        document = docx_mod.Document()
        document.add_paragraph('Nội dung')
        pg_mar = document.sections[0]._sectPr.find(qn('w:pgMar'))
        del pg_mar.attrib[qn('w:top')]

        doc = parse_docx(_blob_of(document))
        self.assertIsNone(doc.pages.margin_mm['top'])
        self.assertIsNotNone(doc.pages.margin_mm['left'],
                             'các lề còn khai vẫn phải đọc được bình thường')

    def test_doan_trong_khong_sap_va_van_co_dinh_dang(self):
        """max() trên list rỗng sẽ nổ nếu thiếu guard; văn bản thật đầy dòng trống."""
        import docx as docx_mod

        document = docx_mod.Document()
        document.add_paragraph('')
        document.add_paragraph('   ')
        document.add_paragraph('Có chữ')

        doc = parse_docx(_blob_of(document))
        self.assertEqual(len(doc.paras), 3)
        self.assertEqual([p.index for p in doc.paras], [0, 1, 2])
        self.assertIsNotNone(doc.paras[0].fmt.size_pt,
                             'đoạn trống vẫn phải lấy được cỡ chữ từ style')
        self.assertFalse(doc.paras[0].runs_conflict)


if __name__ == '__main__':
    unittest.main()
