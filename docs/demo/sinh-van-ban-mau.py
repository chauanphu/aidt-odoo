#!/usr/bin/env python3
"""Sinh văn bản mẫu để trình diễn màn hình kiểm tra thể thức.

Chạy:  python3 docs/demo/sinh-van-ban-mau.py
Kết quả: ba file .docx trong docs/demo/

Sinh bằng code chứ không commit file nhị phân — cùng lý do với fixture của
engine: sửa được, review được bằng diff, và ai cũng dựng lại được.

Ba file phục vụ ba câu chuyện khác nhau trong buổi demo:

  01-dat-chuan.docx      Văn bản đúng thể thức NĐ 30 → engine báo ĐẠT.
                         Cần có, nếu không người xem sẽ nghĩ engine bắt lỗi bừa.

  02-sai-the-thuc.docx   Bốn lỗi cố ý, mỗi lỗi một loại khác nhau: lề trang,
                         chữ đậm, căn lề, thụt đầu dòng. Cho thấy phân biệt
                         được lỗi chặn với cảnh báo.

  03-dau-trang-bang.docx Khối đầu trang dựng bằng bảng 2 cột — đúng cách văn
                         bản hành chính thật được soạn. Trước đây parser không
                         đọc được đoạn nằm trong bảng nên báo "thiếu tiêu đề"
                         cho gần như mọi văn bản thật.
"""
import pathlib

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt

THU_MUC = pathlib.Path(__file__).parent
QUOC_HIEU = 'CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM'
TIEU_DE_DANG = 'ĐẢNG CỘNG SẢN VIỆT NAM'

# Cỡ chữ theo NĐ 30/2020 Phụ lục I. Bộ luật 66-QĐ/TW hiện CHƯA THẨM ĐỊNH nên
# các mẫu này dựng theo chuẩn hành chính.
STYLE_ND30 = {
    'VB_QuocHieu': dict(size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER),
    'VB_SoKyHieu': dict(size=13, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER),
    'VB_TrichYeu': dict(size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER),
    'VB_NoiDung': dict(size=14, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY),
    'VB_NoiNhan': dict(size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT),
    'VB_ChuKy': dict(size=14, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT),
    'VB_TieuDeDang': dict(size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER),
}

THAN_BAI = (
    'Thực hiện chương trình công tác năm 2026, Văn phòng đề nghị các đơn vị '
    'trực thuộc nghiêm túc triển khai những nội dung nêu trên và báo cáo kết '
    'quả về Văn phòng trước ngày 30 tháng 9 năm 2026.'
)


def _tai_lieu(le_mm=(22, 22, 32, 17), sua_style=None):
    tai_lieu = Document()
    khung = tai_lieu.sections[0]
    khung.page_width, khung.page_height = Mm(210), Mm(297)
    tren, duoi, trai, phai = le_mm
    khung.top_margin, khung.bottom_margin = Mm(tren), Mm(duoi)
    khung.left_margin, khung.right_margin = Mm(trai), Mm(phai)

    khai_bao = {ten: dict(gia_tri) for ten, gia_tri in STYLE_ND30.items()}
    for ten, thay_doi in (sua_style or {}).items():
        khai_bao[ten].update(thay_doi)

    for ten, gia_tri in khai_bao.items():
        style = tai_lieu.styles.add_style(ten, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = 'Times New Roman'
        style.font.size = Pt(gia_tri['size'])
        style.font.bold = gia_tri['bold']
        style.paragraph_format.alignment = gia_tri['align']
        if ten == 'VB_NoiDung':
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.first_line_indent = Mm(
                gia_tri.get('thut_mm', 12.7))
    return tai_lieu


def _than_van_ban(tai_lieu, tieu_de_style='VB_QuocHieu', tieu_de=QUOC_HIEU):
    tai_lieu.add_paragraph(tieu_de, style=tieu_de_style)
    tai_lieu.add_paragraph('Số: 145/CV-VPTU', style='VB_SoKyHieu')
    tai_lieu.add_paragraph('Về việc triển khai nhiệm vụ quý III năm 2026',
                           style='VB_TrichYeu')
    tai_lieu.add_paragraph(THAN_BAI, style='VB_NoiDung')
    tai_lieu.add_paragraph('Nơi nhận:', style='VB_NoiNhan')
    tai_lieu.add_paragraph('- Các ban, phòng trực thuộc;', style='VB_NoiNhan')
    tai_lieu.add_paragraph('- Lưu: VT.', style='VB_NoiNhan')
    tai_lieu.add_paragraph('CHÁNH VĂN PHÒNG', style='VB_ChuKy')
    return tai_lieu


def dat_chuan():
    """Đúng thể thức NĐ 30 — engine phải báo đạt, không phát hiện gì."""
    return _than_van_ban(_tai_lieu())


def sai_the_thuc():
    """Bốn lỗi cố ý, mỗi lỗi một loại, để thấy cách phân loại mức nghiêm trọng.

    - lề trên 15mm, ngoài khoảng 20-25mm cho phép
    - trích yếu không in đậm
    - nội dung căn trái thay vì căn đều hai bên
    - thụt đầu dòng 3cm, ngoài khoảng 1-1,27cm
    """
    tai_lieu = _tai_lieu(
        le_mm=(15, 22, 32, 17),
        sua_style={
            'VB_TrichYeu': {'bold': False},
            'VB_NoiDung': {'align': WD_ALIGN_PARAGRAPH.LEFT, 'thut_mm': 30},
        })
    return _than_van_ban(tai_lieu)


def dau_trang_bang():
    """Khối đầu trang dựng bằng bảng 2 cột, như văn bản hành chính thật.

    python-docx không trả đoạn nằm trong bảng qua `document.paragraphs`, nên
    trước bản sửa C1 thì tiêu đề trong bảng biến mất khỏi mắt engine và mọi
    văn bản kiểu này bị báo thiếu tiêu đề.
    """
    tai_lieu = _tai_lieu()
    bang = tai_lieu.add_table(rows=2, cols=2)
    bang.cell(0, 0).paragraphs[0].text = 'VĂN PHÒNG TỈNH ỦY'
    bang.cell(0, 0).paragraphs[0].style = tai_lieu.styles['VB_SoKyHieu']
    o_tieu_de = bang.cell(0, 1).paragraphs[0]
    o_tieu_de.text = QUOC_HIEU
    o_tieu_de.style = tai_lieu.styles['VB_QuocHieu']
    bang.cell(1, 0).paragraphs[0].text = 'Số: 146/CV-VPTU'
    bang.cell(1, 0).paragraphs[0].style = tai_lieu.styles['VB_SoKyHieu']

    tai_lieu.add_paragraph('Về việc báo cáo kết quả thực hiện',
                           style='VB_TrichYeu')
    tai_lieu.add_paragraph(THAN_BAI, style='VB_NoiDung')
    tai_lieu.add_paragraph('Nơi nhận:', style='VB_NoiNhan')
    tai_lieu.add_paragraph('- Lưu: VT.', style='VB_NoiNhan')
    tai_lieu.add_paragraph('CHÁNH VĂN PHÒNG', style='VB_ChuKy')
    return tai_lieu


def main():
    for ten_file, ham in (('01-dat-chuan.docx', dat_chuan),
                          ('02-sai-the-thuc.docx', sai_the_thuc),
                          ('03-dau-trang-bang.docx', dau_trang_bang)):
        duong_dan = THU_MUC / ten_file
        ham().save(duong_dan)
        print('đã sinh %s' % duong_dan)


if __name__ == '__main__':
    main()
